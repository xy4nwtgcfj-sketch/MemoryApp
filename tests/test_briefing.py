import re

import numpy as np
import pytest
from fastapi.testclient import TestClient

from lessons_vault import main
from lessons_vault.briefing import (
    ANPASSEN_PROMPT, BRIEFING_PROMPT, DIALOG_PROMPT, KOMBI_PROMPT, faktor_bestaetigung, faktor_feedback, faktor_schwere,
    kombi_auswerten,
)
from lessons_vault.config import Settings
from lessons_vault.embeddings import HashingProvider
from lessons_vault.main import create_app
from lessons_vault.structuring import Strukturierer

from .test_api import lesson_json
from .test_structuring import FakeClient, payload


KOMBI_ANTWORT = {
    "zusammenfassung": "Hypercare steht und fällt mit klaren Eskalationswegen.",
    "titel": "Hypercare Playbook",
    "untertitel": "Stabilizing the first four weeks",
    "kernbotschaften": [
        {"titel": "Staff before you stabilize", "text": "Confirm shift leads in writing before go-live."},
        {"titel": "", "text": ""},  # leer fällt weg
    ],
    "ausarbeitung": "## Vorbereitung\n- Eskalationsstufen vorab festlegen\n- Schichtleiter schriftlich einplanen [E1]\n- Erfundene Erfahrung [E7]",
    "abschnitte": [
        {"titel": "Vorbereitung", "kernaussage": "Fix staffing and escalation before go-live.", "darstellung": "schritte", "punkte": [
            {"punkt": "Hypercare-Team und Eskalationsstufen festlegen", "quelle": "allgemein", "lesson_nr": 0,
             "kurz": "Escalation path", "icon": "flag", "gewicht": 4, "eintritt": 0, "auswirkung": 0},
            {"punkt": "Schichtleiter für die ersten zwei Wochen schriftlich einplanen", "quelle": "erfahrung", "lesson_nr": 1,
             "kurz": "Staffing confirmed", "icon": "users", "gewicht": 5, "eintritt": 0, "auswirkung": 0},
        ]},
        {"titel": "Betrieb", "darstellung": "quatsch", "punkte": [  # ohne kernaussage, unbekannte Darstellung und Felder
            {"punkt": "Tägliches Hypercare-Meeting", "quelle": "allgemein", "lesson_nr": 0, "icon": "nixda", "gewicht": 9},
            {"punkt": "Behauptete Erfahrung ohne Lesson", "quelle": "erfahrung", "lesson_nr": 0},
        ]},
        {"titel": "Leer", "punkte": []},
    ],
    "naechste_schritte": [
        {"schritt": "Confirm shift leads in writing", "zeitpunkt": "Go-live −2 weeks"},
        {"schritt": "", "zeitpunkt": "nie"},
    ],
    "besonderheiten": [
        {"punkt": "Bei dir fehlte das Personal, nicht die Technik: Personal früher bestätigen als üblich", "lesson_nr": 1},
        {"punkt": "ohne Bezug", "lesson_nr": 0},
    ],
    "luecken": [
        {"titel": "Define exit criteria for hypercare before go-live", "kontext": "Any go-live with a hypercare phase.",
         "wissen": "Exit criteria make the end of hypercare a decision, not a drift. Typical criteria: ticket inflow below a threshold, no open P1/P2, key users confirm.",
         "kern": "Agree measurable hypercare exit criteria before go-live.", "anwendung": "Put the criteria into the cutover plan.",
         "frage": "Have you ever had hypercare drag on because nobody defined its end?", "phase": "hypercare"},
        {"titel": "", "kontext": "", "wissen": "unvollständig", "kern": "", "anwendung": "", "frage": "", "phase": ""},
    ],
}


def antwort(kwargs):
    """Fake-Claude: Briefing-Aufrufe bekommen eine Zusammenfassung, alles andere einen Lesson-Entwurf."""
    if kwargs["system"] == KOMBI_PROMPT:
        return KOMBI_ANTWORT
    if kwargs["system"] == DIALOG_PROMPT:
        return {"antwort": "## Carriers\n- Agree extended pickup windows for the first three days.\n- Your shift leads must be confirmed first [E1]; a fake mark [E9].",
                "anpassung_sinnvoll": True, "anpassung_hinweis": "Add a section on carrier coordination."}
    if kwargs["system"] == ANPASSEN_PROMPT:
        return {**KOMBI_ANTWORT, "titel": "Hypercare Playbook v2",
                "abschnitte": [*KOMBI_ANTWORT["abschnitte"], {"titel": "Carriers", "kernaussage": "Carriers need extended windows.", "darstellung": "kacheln", "punkte": [
                    {"punkt": "Agree extended pickup windows", "quelle": "allgemein", "lesson_nr": 0, "kurz": "Pickup windows", "icon": "calendar", "gewicht": 3}]}],
                "aenderungen": ["Added a Carriers section.", "Kept staffing and escalation unchanged."]}
    if kwargs["system"] == BRIEFING_PROMPT:
        return {
            "zusammenfassung": "Plane die Personalfreigabe früh.",
            "checkliste": [{"punkt": "Schichtleiter schriftlich bestätigen", "lesson_nr": 1}, {"punkt": "", "lesson_nr": 2}],
        }
    return payload()


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "ollama_status", lambda url, timeout=2.0: {"erreichbar": False, "modelle": []})
    claude = FakeClient(antwort)
    settings = Settings(db_path=tmp_path / "t.db", anthropic_api_key="", voyage_api_key="", ollama_model="")
    app = create_app(settings, Strukturierer(claude, "claude-opus-5-5"), HashingProvider())
    client = TestClient(app)

    def add(**kw):
        return client.post("/api/lessons", json=lesson_json(**kw)).json()["lesson"]

    return client, claude, add


def test_faktoren():
    assert faktor_schwere(1) == pytest.approx(0.6) and faktor_schwere(5) == pytest.approx(1.0)
    assert faktor_bestaetigung(1) == 1.0 and faktor_bestaetigung(3) == pytest.approx(1.3) and faktor_bestaetigung(20) == pytest.approx(1.6)
    a = np.array([1.0, 0.0])
    b = np.array([0.0, 1.0])
    assert faktor_feedback("x", [("x", "nicht_relevant", a)], a) == pytest.approx(0.7)
    assert faktor_feedback("x", [("x", "nicht_relevant", b)], a) == pytest.approx(1.0)  # andere Situation zählt nicht
    assert faktor_feedback("x", [("x", "hilfreich", None)], a) == pytest.approx(1.15)
    assert faktor_feedback("x", [("x", "nicht_relevant", a)] * 10, a) == pytest.approx(0.3)


def test_ranking_schwere_und_bestaetigung(env):
    client, _, add = env
    gleich = dict(situation="Projektplanung für eine Rollout-Welle.", ereignis="Zeitplan zu knapp.", bereich="Beruf",
                  lesson="Bei der Projektplanung Puffer für den Cutover einplanen.")
    leicht = add(titel="Projektplanung Puffer A", schweregrad=1, **gleich)
    schwer = add(titel="Projektplanung Puffer B", schweregrad=5, **gleich)
    oft = add(titel="Projektplanung Puffer C", schweregrad=1, **gleich)
    for _ in range(3):
        client.post(f"/api/lessons/{oft['id']}/bestaetigen", json={})
    weg = add(titel="Projektplanung Puffer D", schweregrad=5, **gleich)
    client.put(f"/api/lessons/{weg['id']}", json={"status": "überholt"})

    data = client.post("/api/briefing", json={"situation": "Projektplanung für Rollout Welle 4 schreiben", "modus": "eigen"}).json()
    ids = [t["lesson"]["id"] for t in data["treffer"]]
    assert weg["id"] not in ids
    assert ids.index(schwer["id"]) < ids.index(leicht["id"])
    assert ids.index(oft["id"]) < ids.index(leicht["id"])
    assert len(ids) <= 5
    assert data["checkliste"][0]["punkt"]
    assert data["ki_moeglich"] is True


def test_max_fuenf_und_bereichsfilter(env):
    client, _, add = env
    for i in range(7):
        add(titel=f"Bankgespräch Unterlagen {i}", lesson="Vor dem Bankgespräch Unterlagen prüfen.")
    add(titel="Bankgespräch im Beruf", bereich="Beruf", lesson="Bankgespräch vorbereiten.")
    data = client.post("/api/briefing", json={"situation": "Bankgespräch nächste Woche", "modus": "eigen"}).json()
    assert len(data["treffer"]) == 5
    data = client.post("/api/briefing", json={"situation": "Bankgespräch nächste Woche", "bereich": "Beruf", "modus": "eigen"}).json()
    assert [t["lesson"]["bereich"] for t in data["treffer"]] == ["Beruf"]


def test_nichts_passendes(env):
    client, _, add = env
    add(titel="Notar früh terminieren", lesson="Notartermin früh fixieren.", situation="Kauf.", ereignis="Termin zu spät.", tags=["Notar"])
    data = client.post("/api/briefing", json={"situation": "xyzqq vvvvv", "modus": "eigen"}).json()
    assert data["treffer"] == [] and data["ki_moeglich"] is False


def test_zusammenfassung(env):
    client, claude, add = env
    l1 = add(titel="Cutover Personal bestätigen", bereich="Beruf", lesson="Vor jedem Cutover Personal bestätigen.")
    bid = client.post("/api/briefing", json={"situation": "Cutover planen", "modus": "eigen"}).json()["id"]
    data = client.post(f"/api/briefing/{bid}/zusammenfassung").json()
    assert data["zusammenfassung"] == "Plane die Personalfreigabe früh."
    assert data["ki_checkliste"] == [{"punkt": "Schichtleiter schriftlich bestätigen", "lesson_id": l1["id"]}]
    call = claude.calls[-1]
    assert call["output_config"]["format"]["schema"]["properties"]["checkliste"]["items"]["properties"]["lesson_nr"]["enum"] == [1]
    assert "<situation>\nCutover planen\n</situation>" in call["messages"][0]["content"]
    # Zweiter Aufruf nutzt den gespeicherten Text, kein neuer API-Call
    anzahl = len(claude.calls)
    assert client.post(f"/api/briefing/{bid}/zusammenfassung").json()["zusammenfassung"]
    assert len(claude.calls) == anzahl
    client.post(f"/api/briefing/{bid}/zusammenfassung", params={"neu": True})
    assert len(claude.calls) == anzahl + 1
    # gespeichertes Briefing wieder abrufbar
    assert client.get(f"/api/briefing/{bid}").json()["zusammenfassung"]


def test_lokale_bereiche_nicht_an_claude(env):
    client, claude, add = env
    add(titel="Geheim Cutover Beruf", bereich="Beruf", lesson="Cutover Personal bestätigen.")
    add(titel="Cutover privat Umzug", bereich="Immobilien", lesson="Beim Umzug Cutover der Versorger planen.")
    client.put("/api/einstellungen", json={"strukturierung_bereiche": {"Beruf": "lokal"}})
    data = client.post("/api/briefing", json={"situation": "Cutover planen", "modus": "eigen"}).json()
    assert len(data["treffer"]) == 2 and "lokalen Bereichen" in data["ki_hinweis"]
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    prompt = claude.calls[-1]["messages"][0]["content"]
    assert "Umzug" in prompt and "Geheim" not in prompt

    nur_beruf = client.post("/api/briefing", json={"situation": "Cutover planen", "bereich": "Beruf", "modus": "eigen"}).json()
    assert nur_beruf["ki_moeglich"] is False
    assert client.post(f"/api/briefing/{nur_beruf['id']}/zusammenfassung").status_code == 409


def test_ohne_claude(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "ollama_status", lambda url, timeout=2.0: {"erreichbar": False, "modelle": []})
    app = create_app(Settings(db_path=tmp_path / "t.db", anthropic_api_key="", ollama_model=""), Strukturierer(None, "m"), HashingProvider())
    client = TestClient(app)
    client.post("/api/lessons", json=lesson_json())
    data = client.post("/api/briefing", json={"situation": "Bankgespräch", "modus": "eigen"}).json()
    assert data["treffer"] and data["ki_moeglich"] is False and "ANTHROPIC_API_KEY" in data["ki_hinweis"]


def test_claude_fehler_502(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "ollama_status", lambda url, timeout=2.0: {"erreichbar": False, "modelle": []})
    app = create_app(Settings(db_path=tmp_path / "t.db", anthropic_api_key="", ollama_model=""),
                     Strukturierer(FakeClient(payload(), stop_reason="refusal"), "m"), HashingProvider())
    client = TestClient(app)
    client.post("/api/lessons", json=lesson_json())
    bid = client.post("/api/briefing", json={"situation": "Bankgespräch", "modus": "eigen"}).json()["id"]
    r = client.post(f"/api/briefing/{bid}/zusammenfassung")
    assert r.status_code == 502 and "abgelehnt" in r.json()["detail"]


def test_feedback_verschiebt_ranking(env):
    client, _, add = env
    gleich = dict(situation="Besichtigung einer Wohnung.", ereignis="Mangel übersehen.", bereich="Immobilien", schweregrad=3)
    a = add(titel="Besichtigung Keller prüfen", lesson="Bei jeder Besichtigung den Keller auf Feuchtigkeit prüfen.", **gleich)
    b = add(titel="Besichtigung Dach prüfen", lesson="Bei jeder Besichtigung das Dach auf Schäden prüfen.", **gleich)
    situation = "Besichtigung Wohnung morgen"
    erstes = client.post("/api/briefing", json={"situation": situation, "modus": "eigen"}).json()
    oben = erstes["treffer"][0]["lesson"]["id"]
    unten = b["id"] if oben == a["id"] else a["id"]
    assert client.post(f"/api/briefing/{erstes['id']}/feedback", json={"lesson_id": oben, "bewertung": "nicht_relevant"}).json() == {oben: "nicht_relevant"}
    client.post(f"/api/briefing/{erstes['id']}/feedback", json={"lesson_id": unten, "bewertung": "hilfreich"})
    zweites = client.post("/api/briefing", json={"situation": situation, "modus": "eigen"}).json()
    assert zweites["treffer"][0]["lesson"]["id"] == unten
    assert zweites["treffer"][1]["faktoren"]["feedback"] < 1.0

    # Feedback zurücknehmen
    assert client.post(f"/api/briefing/{erstes['id']}/feedback", json={"lesson_id": oben, "bewertung": ""}).json() == {unten: "hilfreich"}
    stat = client.get("/api/statistik").json()
    assert stat["briefings"] == 2 and stat["briefings_hilfreich"] == 1 and stat["briefings_bewertet"] == 1
    verlauf = client.get("/api/briefings").json()
    assert verlauf[0]["situation"] == situation and verlauf[1]["hilfreich"] == 1
    assert client.delete(f"/api/briefing/{erstes['id']}").status_code == 204
    assert client.get(f"/api/briefing/{erstes['id']}").status_code == 404


def test_geloeschte_lesson_im_alten_briefing(env):
    client, _, add = env
    l = add(titel="Bankgespräch vorbereiten")
    bid = client.post("/api/briefing", json={"situation": "Bankgespräch", "modus": "eigen"}).json()["id"]
    client.delete(f"/api/lessons/{l['id']}")
    assert client.get(f"/api/briefing/{bid}").json()["treffer"] == []


def test_kombi_auswerten_sichert_quellen():
    from lessons_vault.models import Lesson
    l = Lesson(id="L1", titel="t", situation="s", ereignis="e", lesson="r", bereich="Beruf", schweregrad=3,
               status="neu", bestaetigungen=1, erstellt_am="x", geaendert_am="x", quelle="text")
    plan = kombi_auswerten(KOMBI_ANTWORT, [l])
    punkte = [p for a in plan["abschnitte"] for p in a["punkte"]]
    assert [a["titel"] for a in plan["abschnitte"]] == ["Vorbereitung", "Betrieb"]  # leere Abschnitte fallen weg
    assert punkte[1] == {"punkt": "Schichtleiter für die ersten zwei Wochen schriftlich einplanen", "quelle": "erfahrung", "lesson_id": "L1",
                         "kurz": "Staffing confirmed", "icon": "users", "gewicht": 5, "eintritt": 0, "auswirkung": 0,
                         "wert": "", "start": 0, "ende": 0, "aufwand": 0, "gruppe": "", "status": ""}
    # "erfahrung" ohne gültige Lesson wird nie als eigene Erfahrung ausgewiesen
    assert punkte[3]["quelle"] == "allgemein" and punkte[3]["lesson_id"] is None
    # Folienfelder: Abschnitt 1 vollständig, Abschnitt 2 mit Rückfallwerten
    assert plan["abschnitte"][0]["kernaussage"] == "Fix staffing and escalation before go-live." and plan["abschnitte"][0]["darstellung"] == "schritte"
    assert plan["abschnitte"][1]["kernaussage"] == "" and plan["abschnitte"][1]["darstellung"] == "kacheln"
    assert punkte[2] == {"punkt": "Tägliches Hypercare-Meeting", "quelle": "allgemein", "lesson_id": None, "kurz": "", "icon": "", "gewicht": 0,
                         "eintritt": 0, "auswirkung": 0, "wert": "", "start": 0, "ende": 0, "aufwand": 0, "gruppe": "", "status": ""}
    assert plan["titel"] == "Hypercare Playbook" and plan["untertitel"] == "Stabilizing the first four weeks"
    assert plan["kernbotschaften"] == [{"titel": "Staff before you stabilize", "text": "Confirm shift leads in writing before go-live."}]
    assert plan["naechste_schritte"] == [{"schritt": "Confirm shift leads in writing", "zeitpunkt": "Go-live −2 weeks"}]
    assert plan["punkte_erfahrung"] == 1 and plan["punkte_gesamt"] == 4
    assert plan["besonderheiten"] == [{"punkt": KOMBI_ANTWORT["besonderheiten"][0]["punkt"], "lesson_id": "L1"}]
    # Blinde Flecken: nur vollständige Einträge
    assert [l["titel"] for l in plan["luecken"]] == ["Define exit criteria for hypercare before go-live"]
    assert plan["luecken"][0]["phase"] == "hypercare"


def test_rueckschau_bestaetigt_und_erzeugt_entwuerfe(env):
    client, claude, add = env
    l1 = add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
             lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Plan für die Hypercare-Phase eines SAP-EWM-Projekts"}).json()
    r = client.post(f"/api/briefing/{data['id']}/zusammenfassung").json()
    punkte = [p for a in r["plan"]["abschnitte"] for p in a["punkte"]]
    bewertung = [
        {**punkte[0], "ergebnis": "bestaetigt"},                       # Fachwissen, hat gestimmt -> ein Entwurf "bewährt"
        {**punkte[1], "ergebnis": "bestaetigt"},                       # eigener Eintrag -> erneut erlebt
        {**punkte[2], "ergebnis": "abweichung", "notiz": "Meeting fand nur alle zwei Tage statt."},
        {**punkte[3], "ergebnis": ""},
    ]
    entwuerfe_vorher = client.get("/api/statistik").json()["offene_entwuerfe"]
    res = client.post(f"/api/briefing/{data['id']}/rueckschau", json={"punkte": bewertung, "fazit": "Hypercare war kürzer als geplant."})
    assert res.status_code == 200
    body = res.json()
    assert body["bestaetigt"] == 1 and body["entwuerfe"] == 3  # Abweichung, bewährtes Fachwissen, Fazit
    assert body["rueckschau"]["fazit"] == "Hypercare war kürzer als geplant."
    assert client.get(f"/api/lessons/{l1['id']}").json()["lesson"]["bestaetigungen"] == 2
    assert client.get("/api/statistik").json()["offene_entwuerfe"] == entwuerfe_vorher + 3
    assert client.get("/api/statistik").json()["rueckschauen"] == 1
    # Die Notizen sind strukturiert worden (Claude-Aufrufe für Entwürfe) und tragen den Bezug zur Situation
    texte = [e["rohtext"] for e in client.get("/api/entwuerfe").json()["eintraege"]]
    assert any("did not hold up" in t and "alle zwei Tage" in t for t in texte)
    assert any("proved right in practice" in t for t in texte)
    assert any("Looking back at" in t for t in texte)
    # Nur eine Rückschau je Briefing; im Verlauf markiert
    assert client.post(f"/api/briefing/{data['id']}/rueckschau", json={"punkte": [], "fazit": "x"}).status_code == 409
    eintrag = next(b for b in client.get("/api/briefings").json() if b["id"] == data["id"])
    assert eintrag["rueckschau"] is True and eintrag["plan"] is True


def test_luecke_uebernehmen_merkt_sich_den_eintrag(env):
    """Ein blinder Fleck wird einmal übernommen; beim nächsten Öffnen ist die Lücke geschlossen, keine Dublette."""
    client, claude, add = env
    add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
        lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Plan für die Hypercare-Phase eines SAP-EWM-Projekts"}).json()
    plan = client.post(f"/api/briefing/{data['id']}/zusammenfassung").json()["plan"]
    luecke = plan["luecken"][0]
    assert not luecke.get("lesson_id")
    eintrag = {"art": "wissen", "titel": luecke["titel"], "situation": luecke["kontext"], "ereignis": luecke["wissen"],
               "lesson": luecke["kern"], "massnahme": luecke["anwendung"], "bereich": "Beruf", "phase": luecke["phase"],
               "tags": [], "schweregrad": 3, "spezifisch": False, "quelle": "text"}
    vorher = client.get("/api/statistik").json()["gesamt"]
    res = client.post(f"/api/briefing/{data['id']}/luecke/0", json=eintrag)
    assert res.status_code == 201
    lesson_id = res.json()["lesson"]["id"]
    assert client.get(f"/api/lessons/{lesson_id}").json()["lesson"]["art"] == "wissen"
    # Das Briefing kennt den Eintrag; eine zweite Übernahme wird abgelehnt
    assert client.get(f"/api/briefing/{data['id']}").json()["plan"]["luecken"][0]["lesson_id"] == lesson_id
    assert client.post(f"/api/briefing/{data['id']}/luecke/0", json=eintrag).status_code == 409
    assert client.post(f"/api/briefing/{data['id']}/luecke/7", json=eintrag).status_code == 404
    assert client.get("/api/statistik").json()["gesamt"] == vorher + 1
    # Wird der Eintrag gelöscht, ist die Lücke wieder offen und darf erneut übernommen werden
    client.delete(f"/api/lessons/{lesson_id}")
    assert client.get(f"/api/briefing/{data['id']}").json()["plan"]["luecken"][0]["lesson_id"] == ""
    assert client.post(f"/api/briefing/{data['id']}/luecke/0", json=eintrag).status_code == 201


def test_dialog_nachfrage_und_plan_anpassen(env):
    """Nachfragen werden im Dialog beantwortet (mit aufgelösten Marken), Input kann den Plan überarbeiten."""
    client, claude, add = env
    l1 = add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
             lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Plan für die Hypercare-Phase eines SAP-EWM-Projekts"}).json()
    # Ohne Plan keine Nachfrage
    assert client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "Was ist mit Carriern?"}).status_code == 409
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    # Lücke übernehmen, damit die Anpassung sie als geschlossen behalten muss
    luecke = client.get(f"/api/briefing/{data['id']}").json()["plan"]["luecken"][0]
    lid = client.post(f"/api/briefing/{data['id']}/luecke/0", json={
        "art": "wissen", "titel": luecke["titel"], "situation": luecke["kontext"], "ereignis": luecke["wissen"], "lesson": luecke["kern"],
        "massnahme": "", "bereich": "Beruf", "phase": "", "tags": [], "schweregrad": 3, "spezifisch": False, "quelle": "text"}).json()["lesson"]["id"]
    # Nachfrage
    res = client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "What about carriers?", "modus": "frage"})
    assert res.status_code == 200
    dialog = res.json()["dialog"]
    assert [d["rolle"] for d in dialog] == ["nutzer", "claude"] and dialog[0]["art"] == "frage" and dialog[0]["text"] == "What about carriers?"
    assert dialog[1]["art"] == "antwort" and f"[[E:{l1['id']}]]" in dialog[1]["text"] and "[E9]" not in dialog[1]["text"] and "E:None" not in dialog[1]["text"]
    assert dialog[1]["anpassung_sinnvoll"] is True and dialog[1]["anpassung_hinweis"] == "Add a section on carrier coordination."
    assert all(d["zeit"] for d in dialog)
    # Claude sah Aufgabe, Einträge, Plan und Eingabe; der Plan trägt Eintragsnummern statt IDs
    nachricht = claude.calls[-1]["messages"][0]["content"]
    assert "<plan>" in nachricht and "<eingabe>\nWhat about carriers?" in nachricht and "\"lesson_nr\": 1" in nachricht and l1["id"] not in nachricht
    assert res.json()["plan"]["titel"] == "Hypercare Playbook"  # Nachfrage ändert den Plan nicht
    # Anpassen: neuer Plan, Änderungen im Dialog, übernommene Lücke bleibt geschlossen, Dialog ging mit
    res = client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "Add carriers", "modus": "anpassen"})
    assert res.status_code == 200
    body = res.json()
    assert body["plan"]["titel"] == "Hypercare Playbook v2" and [a["titel"] for a in body["plan"]["abschnitte"]][-1] == "Carriers"
    assert body["plan"]["luecken"][0]["lesson_id"] == lid
    assert [d["art"] for d in body["dialog"]] == ["frage", "antwort", "anpassen", "anpassung"]
    assert "Added a Carriers section." in body["dialog"][-1]["text"]
    assert "<dialog>" in claude.calls[-1]["messages"][0]["content"] and "Nutzer: What about carriers?" in claude.calls[-1]["messages"][0]["content"]
    # PowerPoint nutzt den angepassten Plan
    from pptx import Presentation
    import io
    praes = Presentation(io.BytesIO(client.get(f"/api/briefing/{data['id']}/pptx").content))
    assert any("Carriers" in sh.text_frame.text for s in praes.slides for sh in s.shapes if sh.has_text_frame)
    # Sperrliste gilt auch im Dialog; Moduswechsel verwirft den Dialog; nach der Rückschau keine Anpassung mehr
    client.put("/api/einstellungen", json={"sperrliste": ["Geheimkunde"]})
    assert client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "Geheimkunde will more"}).status_code == 422
    punkte = [p for a in body["plan"]["abschnitte"] for p in a["punkte"]]
    client.post(f"/api/briefing/{data['id']}/rueckschau", json={"punkte": [{**punkte[0], "ergebnis": "bestaetigt"}], "fazit": "ok"})
    assert client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "Change it", "modus": "anpassen"}).status_code == 409
    assert client.post(f"/api/briefing/{data['id']}/dialog", json={"text": "Still a question?", "modus": "frage"}).status_code == 200
    client.post(f"/api/briefing/{data['id']}/zusammenfassung?modus=eigen")
    assert client.get(f"/api/briefing/{data['id']}").json()["dialog"] == []


def test_marke_und_logo_fuer_powerpoint(env):
    import io

    from PIL import Image
    from pptx import Presentation

    client, claude, add = env
    # Voreinstellung: Beruf trägt die Deloitte-Marke, Immobilien van Newboert
    einst = client.get("/api/einstellungen").json()
    assert einst["marken"]["Beruf"]["name"] == "Deloitte" and einst["marken"]["Beruf"]["akzent"] == "#86BC25"
    assert einst["marken"]["Immobilien"]["name"] == "VAN NEWBOERT" and einst["marken"]["Immobilien"]["stil"] == "hell"
    # Marke anpassen, unbekannte Bereiche fallen weg
    r = client.put("/api/einstellungen", json={"marken": {**einst["marken"], "Beruf": {**einst["marken"]["Beruf"], "untertitel": "Internal"}, "Nix": {"name": "x"}}}).json()
    assert r["marken"]["Beruf"]["untertitel"] == "Internal" and "Nix" not in r["marken"]
    assert client.put("/api/einstellungen", json={"marken": {"Beruf": {"farbe": "rot"}}}).status_code == 422
    # Logo hochladen (JPEG wird zu PNG), abrufen, in der Übersicht markiert
    bild = io.BytesIO()
    Image.new("RGB", (300, 120), (0, 0, 0)).save(bild, "JPEG")
    assert client.get("/api/marken/Beruf/logo").status_code == 404
    assert client.post("/api/marken/Beruf/logo", files={"datei": ("logo.jpg", bild.getvalue(), "image/jpeg")}).status_code == 201
    assert client.post("/api/marken/Beruf/logo", files={"datei": ("logo.txt", b"kein bild", "text/plain")}).status_code == 422
    assert client.post("/api/marken/Fremd/logo", files={"datei": ("logo.jpg", bild.getvalue(), "image/jpeg")}).status_code == 404
    logo = client.get("/api/marken/Beruf/logo")
    assert logo.status_code == 200 and logo.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert client.get("/api/marken").json()["Beruf"]["logo"] is True

    add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
        lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Hypercare-Plan"}).json()  # ohne Bereich: Marke des häufigsten Treffer-Bereichs
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    praes = Presentation(io.BytesIO(client.get(f"/api/briefing/{data['id']}/pptx").content))
    texte = [" ".join(sh.text_frame.text for sh in folie.shapes if sh.has_text_frame) for folie in praes.slides]
    assert "Deloitte" in texte[0] and "Internal" in texte[0]
    assert any("Further topics" in t and "Define exit criteria" in t for t in texte)
    bilder = [sh for folie in praes.slides for sh in folie.shapes if sh.shape_type == 13]  # PICTURE
    assert len(bilder) >= len(praes.slides) - 1  # Logo auf jeder Folie
    # Logo entfernen
    assert client.delete("/api/marken/Beruf/logo").status_code == 204
    assert client.get("/api/marken").json()["Beruf"]["logo"] is False

    # Heller Stil (van Newboert): Wortmarke auf den Inhaltsfolien, Illustration nur auf der Titelfolie
    add(titel="WEG-Protokolle vor dem Kauf lesen", bereich="Immobilien", lesson="Vor dem Kauf die WEG-Protokolle der letzten drei Jahre lesen.")
    assert client.post("/api/marken/Immobilien/logo", files={"datei": ("logo.jpg", bild.getvalue(), "image/jpeg")}).status_code == 201
    data = client.post("/api/briefing", json={"situation": "Besichtigung einer Eigentumswohnung", "bereich": "Immobilien"}).json()
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    praes = Presentation(io.BytesIO(client.get(f"/api/briefing/{data['id']}/pptx").content))
    texte = [" ".join(sh.text_frame.text for sh in folie.shapes if sh.has_text_frame) for folie in praes.slides]
    assert "VAN NEWBOERT" in texte[0] and "IMMOBILIEN" in texte[0]
    assert all("VAN NEWBOERT" in t for t in texte[1:])  # Wortmarke oben rechts
    bilder = [sh for folie in praes.slides for sh in folie.shapes if sh.shape_type == 13]
    assert len(bilder) == 1  # Illustration nur auf der Titelfolie
    schriften = {r.font.name for sh in praes.slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert {"Georgia", "Calibri"} <= schriften


def test_kombiniert_ist_standard(env):
    client, claude, add = env
    l1 = add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
             lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Plan für die Hypercare-Phase eines SAP-EWM-Projekts"}).json()
    assert data["modus"] == "kombiniert" and data["ki_moeglich"]
    r = client.post(f"/api/briefing/{data['id']}/zusammenfassung").json()
    call = claude.calls[-1]
    assert call["system"] == KOMBI_PROMPT
    assert call["output_config"]["effort"] == "medium" and call["max_tokens"] == 16000
    assert call["output_config"]["format"]["schema"]["properties"]["abschnitte"]["items"]["properties"]["punkte"]["items"]["properties"]["lesson_nr"]["enum"] == [0, 1]
    assert r["plan"]["abschnitte"][0]["punkte"][1]["lesson_id"] == l1["id"]
    assert r["zusammenfassung"].startswith("Hypercare")
    assert {p["quelle"] for p in r["ki_checkliste"]} == {"erfahrung", "allgemein"}
    stat = client.get("/api/statistik").json()
    assert stat["plan_punkte_erfahrung"] == 1 and stat["plan_punkte_gesamt"] == 4


def test_kombiniert_ohne_eigene_lessons(env):
    client, claude, _ = env
    data = client.post("/api/briefing", json={"situation": "Hypercare planen"}).json()
    assert data["treffer"] == [] and data["ki_moeglich"] is True
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    assert "keine passenden vorhanden" in claude.calls[-1]["messages"][0]["content"]
    eigen = client.post("/api/briefing", json={"situation": "Hypercare planen", "modus": "eigen"}).json()
    assert eigen["ki_moeglich"] is False


def test_moduswechsel(env):
    client, claude, add = env
    add(titel="Hypercare Schichtleiter", lesson="In der Hypercare Schichtleiter einplanen.")
    bid = client.post("/api/briefing", json={"situation": "Hypercare planen"}).json()["id"]
    r = client.post(f"/api/briefing/{bid}/zusammenfassung").json()
    assert r["plan"] is not None
    r = client.post(f"/api/briefing/{bid}/zusammenfassung", params={"modus": "eigen"}).json()
    assert r["modus"] == "eigen" and r["plan"] is None and claude.calls[-1]["system"] == BRIEFING_PROMPT
    anzahl = len(claude.calls)
    client.post(f"/api/briefing/{bid}/zusammenfassung", params={"modus": "eigen"})  # gleicher Modus: gespeichert
    assert len(claude.calls) == anzahl
    assert client.post(f"/api/briefing/{bid}/zusammenfassung", params={"modus": "quatsch"}).status_code == 422


def test_kombiniert_lokal_bereich_nicht_an_claude(env):
    client, claude, add = env
    add(titel="Geheim Hypercare Beruf", bereich="Beruf", lesson="Hypercare Personal bestätigen.")
    client.put("/api/einstellungen", json={"strukturierung_bereiche": {"Beruf": "lokal"}})
    data = client.post("/api/briefing", json={"situation": "Hypercare planen"}).json()
    client.post(f"/api/briefing/{data['id']}/zusammenfassung")
    assert "Geheim" not in claude.calls[-1]["messages"][0]["content"]
    nur_beruf = client.post("/api/briefing", json={"situation": "Hypercare planen", "bereich": "Beruf"}).json()
    assert nur_beruf["ki_moeglich"] is False


def test_ausarbeitung_marken_und_powerpoint(env):
    import io

    from pptx import Presentation

    client, claude, add = env
    l1 = add(titel="Schichtleiter in der Hypercare einplanen", bereich="Beruf",
             lesson="In der Hypercare die Schichtleiter für zwei Wochen schriftlich einplanen.")
    data = client.post("/api/briefing", json={"situation": "Plan für die Hypercare-Phase eines SAP-EWM-Projekts"}).json()
    # Vor dem Plan gibt es keine PowerPoint
    assert client.get(f"/api/briefing/{data['id']}/pptx").status_code == 409
    r = client.post(f"/api/briefing/{data['id']}/zusammenfassung").json()
    text = r["plan"]["ausarbeitung"]
    assert f"[[E:{l1['id']}]]" in text  # gültige Marke aufgelöst
    assert "[E7]" not in text and "E:" not in text.split("Erfundene Erfahrung")[1]  # ungültige Marke entfernt
    assert "ausarbeitung" in claude.calls[-1]["output_config"]["format"]["schema"]["required"]

    res = client.get(f"/api/briefing/{data['id']}/pptx")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    assert "filename*=UTF-8''Briefing%20Plan" in res.headers["content-disposition"]
    praes = Presentation(io.BytesIO(res.content))
    texte = [" ".join(sh.text_frame.text for sh in folie.shapes if sh.has_text_frame) for folie in praes.slides]
    assert "Hypercare Playbook" in texte[0] and "Plan für die Hypercare-Phase" not in texte[0]
    assert "Agenda" in texte[1] and "Vorbereitung" in texte[1] and "Next steps" in texte[1]
    assert any("Schichtleiter für die ersten zwei Wochen" in t for t in texte)
    # Keine Quellen, keine KI, keine Lesson-Titel auf den Folien
    gesamt = " ".join(texte)
    assert "Schichtleiter in der Hypercare einplanen" not in gesamt
    assert not re.search(r"Claude|own experience|knowledge base|★", gesamt)
    # Der ausführliche Text steht in den Notizen der Executive Summary, ohne interne Marken
    notizen = praes.slides[2].notes_slide.notes_text_frame.text
    assert "Schichtleiter schriftlich einplanen" in notizen and "own experience" not in notizen and "[[E:" not in notizen
