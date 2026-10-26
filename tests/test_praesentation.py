"""PowerPoint im Beratungsstil: Aufbau, Darstellungen, Marken und Rückfall für alte Pläne ohne Folienfelder."""
import io
import re

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lessons_vault.models import STANDARD_MARKEN, Marke
from lessons_vault.praesentation import PIKTOGRAMME, briefing_pptx, dateiname, text_ohne_marken
from lessons_vault.briefing import ICONS


def punkt(text, kurz="", icon="", gewicht=0, eintritt=0, auswirkung=0, quelle="allgemein"):
    return {"punkt": text, "quelle": quelle, "lesson_id": "L1" if quelle == "erfahrung" else None, "kurz": kurz, "icon": icon,
            "gewicht": gewicht, "eintritt": eintritt, "auswirkung": auswirkung}


PLAN = {
    "zusammenfassung": "Staffing first. Then escalation.",
    "titel": "Hypercare Playbook: Warehouse Go-live",
    "untertitel": "How to stabilize operations in the first four weeks",
    "kernbotschaften": [
        {"titel": "Staff before you stabilize", "text": "Shift leads must be confirmed before go-live."},
        {"titel": "Make escalation boring", "text": "Three levels with named owners."},
        {"titel": "Define the end on day one", "text": "Measurable exit criteria."},
    ],
    "ausarbeitung": "## Before go-live\n- Confirm staffing [[E:L1]]\n- Agree escalation levels",
    "abschnitte": [
        {"titel": "Preparation", "kernaussage": "Fix staffing before go-live.", "darstellung": "kacheln", "punkte": [
            punkt("Confirm shift leads in writing.", "Staffing confirmed", "users", 5, quelle="erfahrung"),
            punkt("Define three escalation levels.", "Escalation path", "flag", 4),
            punkt("Agree exit criteria.", "Exit criteria", "target", 4),
            punkt("Prepare the war room.", "War room", "gear", 3),
        ]},
        {"titel": "Daily operation", "kernaussage": "A fixed cadence keeps issues visible.", "darstellung": "schritte", "punkte": [
            punkt("Morning stand-up.", "Stand-up", "clock", 4), punkt("Floor walk.", "Floor walk", "search", 4),
            punkt("Cut-off checkpoint.", "Checkpoint", "chart", 3),
        ]},
        {"titel": "Risks", "kernaussage": "Staffing gaps are the main risk.", "darstellung": "risiken", "punkte": [
            punkt("Shift leads unavailable.", "Staffing gap", "users", 5, 3, 3), punkt("Printing fails.", "Printing", "warning", 4, 2, 3),
            punkt("Master data errors.", "Master data", "document", 3, 2, 2),
        ]},
        {"titel": "Priorities", "kernaussage": "Outbound first.", "darstellung": "prioritaeten", "punkte": [
            punkt("Keep the outbound cut-off.", "Outbound", "target", 5), punkt("Resolve P1 within the shift.", "P1", "warning", 4),
            punkt("Collect reporting wishes.", "Reporting", "list", 1),
        ]},
        {"titel": "Handover", "kernaussage": "Transfer to run in a controlled step.", "darstellung": "schritte", "punkte": [
            punkt(f"Step {i}.", f"Step {i}", "check", 2) for i in range(1, 7)
        ]},
    ],
    "besonderheiten": [{"punkt": "Staffing, not technology, was the bottleneck.", "lesson_id": "L1"}],
    "luecken": [{"titel": "Define exit criteria before go-live", "kontext": "c", "wissen": "Exit criteria make the end a decision.",
                 "kern": "Agree measurable exit criteria.", "anwendung": "Put them into the cutover plan.", "frage": "", "phase": "hypercare"}],
    "naechste_schritte": [{"schritt": f"Step {i}", "zeitpunkt": f"Week {i}"} for i in range(1, 6)],
    "punkte_erfahrung": 1, "punkte_gesamt": 19,
}


def texte_von(praes):
    return [" ".join(sh.text_frame.text for sh in folie.shapes if sh.has_text_frame) for folie in praes.slides]


def typen(folie):
    return [sh.shape_type for sh in folie.shapes]


def test_aufbau_beratungsdeck():
    praes = Presentation(io.BytesIO(briefing_pptx("Hypercare plan", PLAN, {}, "2026-10-26", marke=STANDARD_MARKEN["Beruf"])))
    texte = texte_von(praes)
    # Titel, Agenda, Executive Summary, Fokus, je Abschnitt Kapitel + Inhalt (5x2), Watch-outs, Further topics, Next steps
    assert len(praes.slides) == 4 + 10 + 3
    # Titelfolie zeigt den Decktitel, nicht die Aufgabenbeschreibung; die Fußzeile Marke und Titel
    assert "Hypercare Playbook" in texte[0] and "first four weeks" in texte[0] and "Deloitte" in texte[0]
    assert "Hypercare plan" not in texte[0] and "BRIEFING" not in texte[0]
    assert "Agenda" in texte[1] and "Deloitte  ·  Hypercare Playbook" in texte[1] and "Briefing" not in texte[1]
    for kapitel in ("01", "Executive summary", "02", "Preparation", "06", "Handover", "07", "Watch-outs", "08", "Further topics", "09", "Next steps"):
        assert kapitel in texte[1]
    # Executive Summary: drei Kernbotschaften, Storyline, Notizen mit der Ausarbeitung ohne Marken
    assert "What matters most" in texte[2] and "Staff before you stabilize" in texte[2] and "STORYLINE" in texte[2] and "02  Preparation" in texte[2]
    notizen = praes.slides[2].notes_slide.notes_text_frame.text
    assert "Confirm staffing" in notizen and "[[E:" not in notizen and "own experience" not in notizen
    # Fokus: natives Balkendiagramm mit einem Balken je Abschnitt
    assert "Where to focus" in texte[3] and MSO_SHAPE_TYPE.CHART in typen(praes.slides[3])
    diagramm = next(sh for sh in praes.slides[3].shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART).chart
    assert list(diagramm.plots[0].categories) == ["02 Preparation", "03 Daily operation", "04 Risks", "05 Priorities", "06 Handover"]
    assert list(diagramm.plots[0].series[0].values) == [4.0, 3.7, 4.0, 3.3, 2.0]
    # Kapitelfolie mit Nummer und Kernaussage, Inhaltsfolie mit Kernaussage als Titel und Piktogrammen (Freiformen)
    assert "02" in texte[4] and "Preparation" in texte[4] and "Fix staffing before go-live." in texte[4]
    assert "Fix staffing before go-live." in texte[5] and "Staffing confirmed" in texte[5]
    assert typen(praes.slides[5]).count(MSO_SHAPE_TYPE.FREEFORM) == 4
    # Schritte als Chevrons, Risiken als Matrix mit Achsen, Prioritäten mit Diagramm, 6 Schritte als Zeitstrahl
    assert "1  Stand-up" in texte[7] and "IMPACT" in texte[9] and "LIKELIHOOD" in texte[9]
    assert MSO_SHAPE_TYPE.CHART in typen(praes.slides[11]) and "IMPORTANCE" in texte[11]
    prio = next(sh for sh in praes.slides[11].shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART).chart
    assert list(prio.plots[0].categories) == ["Outbound", "P1", "Reporting"]
    assert "Step 6" in texte[13]
    assert "Watch-outs" in texte[14] and "bottleneck" in texte[14]
    assert "Further topics" in texte[15] and "Define exit criteria before go-live" in texte[15]
    assert "Exit criteria make the end a decision" in praes.slides[15].notes_slide.notes_text_frame.text
    assert "Next steps" in texte[16] and "Week 5" in texte[16]
    # Keine Hinweise auf KI, Quellen oder eigene Erfahrung; Seitenzahlen auf den Inhaltsfolien
    gesamt = " ".join(texte)
    assert not re.search(r"Claude|own experience|knowledge base|★|verified", gesamt)
    assert texte[1].rstrip().endswith("2") and texte[16].rstrip().endswith("17")


def test_alter_plan_ohne_folienfelder():
    alt = {
        "zusammenfassung": "Besetzung früh klären. Eskalation festlegen. Exit definieren. Vierter Satz.",
        "ausarbeitung": "x",
        "abschnitte": [
            {"titel": "Vorbereitung", "punkte": [{"punkt": "Team festlegen", "quelle": "allgemein", "lesson_id": None}] * 7},
            {"titel": "Betrieb", "punkte": [{"punkt": "Tägliches Meeting", "quelle": "allgemein", "lesson_id": None}]},
        ],
        "besonderheiten": [], "luecken": [], "punkte_erfahrung": 0, "punkte_gesamt": 8,
    }
    praes = Presentation(io.BytesIO(briefing_pptx("Hypercare planen", alt, {}, marke=None)))
    texte = texte_von(praes)
    assert "Hypercare planen" in texte[0]  # ohne Decktitel bleibt die Situation die Überschrift
    # Kein Fokus-Diagramm ohne Gewichte; 7 Punkte werden auf zwei Kachelfolien verteilt
    assert len(praes.slides) == 3 + 3 + 2
    assert not any(MSO_SHAPE_TYPE.CHART in typen(f) for f in praes.slides)
    assert "Key message 1" in texte[2] and "Besetzung früh klären." in texte[2] and "Vierter Satz" not in texte[2]
    assert "(1/2)" in texte[4] and "(2/2)" in texte[5]
    assert "Vorbereitung" in texte[4]  # ohne Kernaussage ist der Abschnittstitel die Überschrift


def test_helle_marke_und_piktogramme():
    marke = STANDARD_MARKEN["Immobilien"]
    praes = Presentation(io.BytesIO(briefing_pptx("Besichtigung", PLAN, {}, marke=marke)))
    texte = texte_von(praes)
    assert all("VAN NEWBOERT" in t for t in texte[1:])
    schriften = {r.font.name for sh in praes.slides[4].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert "Georgia" in schriften
    # Alle Piktogramme des Schemas sind gezeichnet
    assert set(ICONS) <= set(PIKTOGRAMME)
    # Langer Titel wird kleiner gesetzt statt zu überlaufen
    lang = {**PLAN, "abschnitte": [{**PLAN["abschnitte"][0], "kernaussage": "Fix staffing, escalation paths and measurable exit criteria " * 2}]}
    praes = Presentation(io.BytesIO(briefing_pptx("x", lang, {}, marke=marke)))
    groessen = [r.font.size.pt for sh in praes.slides[5].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs
                if "Fix staffing" in r.text]
    assert groessen and groessen[0] < 26


def pkt2(text, kurz, gewicht=3, **extra):
    return {**punkt(text, kurz, "flag", gewicht), **extra}


def test_alle_zwanzig_darstellungen_und_abwechslung():
    from lessons_vault.briefing import DARSTELLUNGEN
    from lessons_vault.praesentation import BAUSTEINE, darstellung_waehlen

    assert len(DARSTELLUNGEN) == 20 and set(DARSTELLUNGEN) - {"schritte"} <= set(BAUSTEINE)
    basis = [pkt2(f"Point {i} with some text.", f"P{i}", gewicht=(i % 5) + 1, aufwand=(i % 3) + 1, start=i + 1, ende=i + 2,
                  wert=f"{i} d", eintritt=(i % 3) + 1, auswirkung=3, status=("gruen", "gelb", "rot")[i % 3],
                  gruppe=("strength", "weakness", "opportunity", "threat")[i % 4]) for i in range(4)]
    abschnitte = [{"titel": d, "kernaussage": f"Message {d}.", "darstellung": d, "punkte": basis} for d in DARSTELLUNGEN]
    plan = {**PLAN, "abschnitte": abschnitte}
    praes = Presentation(io.BytesIO(briefing_pptx("All layouts", plan, {}, marke=STANDARD_MARKEN["Beruf"])))
    texte = texte_von(praes)
    # Jede Darstellung bekommt eine Kapitel- und eine Inhaltsfolie; Diagramme und Tabellen sind nativ
    assert len(praes.slides) == 4 + 2 * 20 + 3
    typen_gesamt = [t for f in praes.slides for t in typen(f)]
    assert typen_gesamt.count(MSO_SHAPE_TYPE.CHART) == 1 + 4  # Fokus + prioritaeten, radar, kreis, saeulen
    assert MSO_SHAPE_TYPE.TABLE in typen_gesamt
    for erwartet in ("QUICK WINS", "STRENGTHS", "ON TRACK", "W1", "IMPACT", "LIKELIHOOD", "WEAKNESS", "Message pyramide."):
        assert any(erwartet in t for t in texte), erwartet
    # Wiederholungen weichen aus: dreimal kacheln -> kacheln, nummern, tabelle; zweimal prioritaeten -> kreis
    benutzt: dict[str, int] = {}
    k = {"darstellung": "kacheln", "punkte": basis}
    assert [darstellung_waehlen(k, benutzt) for _ in range(3)] == ["kacheln", "nummern", "tabelle"]
    p = {"darstellung": "prioritaeten", "punkte": basis}
    assert [darstellung_waehlen(p, benutzt) for _ in range(2)] == ["prioritaeten", "kreis"]
    # Fehlende Daten: Darstellung fällt zurück
    ohne = [punkt("x", "x"), punkt("y", "y"), punkt("z", "z")]
    assert darstellung_waehlen({"darstellung": "matrix", "punkte": ohne}, {}) == "nummern"
    assert darstellung_waehlen({"darstellung": "swot", "punkte": ohne}, {}) == "kacheln"
    assert darstellung_waehlen({"darstellung": "zeitplan", "punkte": ohne}, {}) == "schritte_v"
    assert darstellung_waehlen({"darstellung": "schritte", "punkte": ohne}, {}) == "schritte_h"
    assert darstellung_waehlen({"darstellung": "kreislauf", "punkte": ohne[:2]}, {}) == "schritte_v"


def test_hilfsfunktionen():
    assert text_ohne_marken("Confirm staffing [[E:abc-1]] early") == "Confirm staffing early"
    assert dateiname("Plan: Hypercare / Welle 4?") == "Briefing Plan Hypercare  Welle 4.pptx"
    assert Marke(name="x").stil == "dunkel"
    with pytest.raises(Exception):
        Marke(farbe="rot")
